"""Regenerate XLSX demo files from the checked-in CSV sources."""
from csv import reader
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

root=Path(__file__).parent
for csv_path in root.glob("sample_*.csv"):
    workbook=Workbook(); sheet=workbook.active; sheet.title="Blood Test"
    with csv_path.open(encoding="utf-8-sig",newline="") as source:
        for row in reader(source): sheet.append(row)
    for cell in sheet[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="0D746F")
    sheet.freeze_panes="A2"
    for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=max(12,max(len(str(c.value or "")) for c in column)+2)
    workbook.save(csv_path.with_suffix(".xlsx"))
