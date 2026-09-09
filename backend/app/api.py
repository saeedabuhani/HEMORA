from datetime import date
from typing import Annotated
import jwt
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from io import BytesIO
import os
from bidi.algorithm import get_display
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session
from .database import get_db
from . import models as m
from .schemas import BloodTestIn, DoctorAssignment, LoginIn, PatientIn, TokenPair
from .security import create_token, decode_token, decrypt_national_id, mask_national_id, national_id_hash, verify_password
from .services import AuditService, ClinicalAnalysisEngine, DomainError, ExplanationService, ImportService, LongitudinalTrendService, PanelCompletenessService, PatientService, RecommendationEngine, TestComparisonEngine
from .limiting import limiter
from .reporting import ReportService

router=APIRouter(prefix="/api"); bearer=HTTPBearer(auto_error=False)
FORBIDDEN = {"code":"FORBIDDEN","message":"אין הרשאה לצפות בנתוני מטופל זה","details":{}}

def scope_patients(query, user):
    """Restrict any Patient query to what this role may see.

    ADMIN sees every patient, a DOCTOR only patients assigned to them,
    a CLINIC only patients belonging to it, and a PATIENT only themselves.
    """
    if user.role == m.Role.ADMIN:
        return query
    if user.role == m.Role.DOCTOR:
        assigned = select(m.DoctorPatient.patient_id).where(m.DoctorPatient.doctor_user_id == user.id)
        return query.where(m.Patient.id.in_(assigned))
    if user.role == m.Role.CLINIC:
        return query.where(m.Patient.clinic_id == user.clinic_id)
    return query.where(m.Patient.id == user.patient_id)

def assert_patient_access(db, user, patient_id):
    """Return the patient, or raise 404 when absent and 403 when out of scope."""
    patient = db.get(m.Patient, patient_id)
    if not patient:
        raise HTTPException(404, detail={"code":"NOT_FOUND","message":"המטופל לא נמצא","details":{}})
    visible = db.scalar(scope_patients(select(m.Patient.id).where(m.Patient.id == patient_id), user))
    if not visible:
        raise HTTPException(403, detail=FORBIDDEN)
    return patient

def authorized_test(db, user, test_id):
    test = db.get(m.BloodTest, test_id)
    if not test:
        raise HTTPException(404, detail={"code":"NOT_FOUND","message":"הבדיקה לא נמצאה","details":{}})
    assert_patient_access(db, user, test.patient_id)
    return test
def current_user(credentials:Annotated[HTTPAuthorizationCredentials|None,Depends(bearer)],db:Session=Depends(get_db)):
    if not credentials: raise HTTPException(401,detail={"code":"NOT_AUTHENTICATED","message":"נדרשת התחברות","details":{}})
    try: payload=decode_token(credentials.credentials)
    except jwt.PyJWTError: raise HTTPException(401,detail={"code":"INVALID_TOKEN","message":"פג תוקף ההתחברות","details":{}})
    user=db.get(m.User,int(payload["sub"]))
    if not user or not user.is_active: raise HTTPException(401,detail={"code":"INACTIVE_USER","message":"החשבון אינו פעיל","details":{}})
    return user
def roles(*allowed):
    def dependency(user=Depends(current_user)):
        if user.role not in allowed: raise HTTPException(403,detail={"code":"FORBIDDEN","message":"אין הרשאה לפעולה זו","details":{}})
        return user
    return dependency
def patient_dict(p,db):
    dates=db.scalars(select(m.BloodTest.test_date).where(m.BloodTest.patient_id==p.id).order_by(m.BloodTest.test_date.desc())).all()
    return {"id":p.id,"first_name":p.first_name,"last_name":p.last_name,"masked_national_id":mask_national_id(decrypt_national_id(p.national_id_encrypted)),"date_of_birth":p.date_of_birth,"biological_sex":p.biological_sex,"phone":p.phone,"email":p.email,"demo":p.demo,"clinic_id":p.clinic_id,"clinic_name":(db.get(m.Clinic,p.clinic_id).name if p.clinic_id else None),"test_count":len(dates),"last_test_date":dates[0] if dates else None}

@router.post("/auth/login",response_model=TokenPair)
@limiter.limit("10/minute")
def login(data:LoginIn,request:Request,db:Session=Depends(get_db)):
    user=db.scalar(select(m.User).where(func.lower(m.User.email)==data.email.lower()))
    if not user or not verify_password(data.password,user.password_hash): raise HTTPException(401,detail={"code":"INVALID_CREDENTIALS","message":"פרטי ההתחברות שגויים","details":{}})
    AuditService.record(db,user.id,"LOGIN","user",user.id,{"ip":request.client.host if request.client else None}); db.commit()
    return TokenPair(access_token=create_token(user.id,user.role.value),refresh_token=create_token(user.id,user.role.value,"refresh"),role=user.role)
@router.post("/auth/refresh")
def refresh(refresh_token:str,db:Session=Depends(get_db)):
    try: payload=decode_token(refresh_token,"refresh")
    except jwt.PyJWTError: raise HTTPException(401,detail={"code":"INVALID_REFRESH","message":"נדרשת התחברות מחדש","details":{}})
    user=db.get(m.User,int(payload["sub"])); return {"access_token":create_token(user.id,user.role.value),"token_type":"bearer"}
@router.get("/auth/me")
def me(user=Depends(current_user)): return {"id":user.id,"email":user.email,"role":user.role}

@router.get("/patients")
def patients(search:str="",page:int=1,page_size:int=20,user=Depends(current_user),db:Session=Depends(get_db)):
    q=scope_patients(select(m.Patient),user)
    if search and user.role!=m.Role.PATIENT:
        terms=[m.Patient.first_name.ilike(f"%{search}%"),m.Patient.last_name.ilike(f"%{search}%")]
        if search.isdigit(): terms.append(m.Patient.national_id_hash==national_id_hash(search))
        accession_ids=select(m.BloodTest.patient_id).where(m.BloodTest.accession_number.ilike(f"%{search}%")); terms.append(m.Patient.id.in_(accession_ids)); q=q.where(or_(*terms))
    total=db.scalar(select(func.count()).select_from(q.subquery())); items=db.scalars(q.order_by(m.Patient.last_name).offset((page-1)*page_size).limit(page_size)).all()
    return {"items":[patient_dict(p,db) for p in items],"total":total,"page":page,"page_size":page_size}
@router.post("/patients")
def create_patient(data:PatientIn,user=Depends(roles(m.Role.DOCTOR,m.Role.CLINIC,m.Role.ADMIN)),db:Session=Depends(get_db)):
    clinic_id=data.clinic_id if user.role==m.Role.ADMIN else user.clinic_id
    if clinic_id is None: raise HTTPException(422,detail={"code":"CLINIC_REQUIRED","message":"יש לשייך את המטופל למרפאה","details":{}})
    if not db.get(m.Clinic,clinic_id): raise HTTPException(422,detail={"code":"UNKNOWN_CLINIC","message":"המרפאה אינה קיימת","details":{}})
    try: p=PatientService.create(db,data,user.id,clinic_id)
    except DomainError as e: raise HTTPException(e.status,detail={"code":e.code,"message":e.message,"details":e.details})
    if user.role==m.Role.DOCTOR:
        db.add(m.DoctorPatient(doctor_user_id=user.id,patient_id=p.id,assigned_by=user.id))
        AuditService.record(db,user.id,"DOCTOR_ASSIGNED","patient",p.id,{"doctor_user_id":user.id}); db.commit()
    return patient_dict(p,db)
@router.get("/patients/{patient_id}")
def patient(patient_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    p=assert_patient_access(db,user,patient_id)
    AuditService.record(db,user.id,"PATIENT_VIEWED","patient",p.id); db.commit(); return patient_dict(p,db)
@router.get("/patients/{patient_id}/tests")
def patient_tests(patient_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    assert_patient_access(db,user,patient_id)
    tests=db.scalars(select(m.BloodTest).where(m.BloodTest.patient_id==patient_id).order_by(m.BloodTest.test_date.desc())).all()
    return [{"id":t.id,"test_date":t.test_date,"accession_number":t.accession_number,"panel":t.panel,"status":t.status,"results_count":len(t.results)} for t in tests]

@router.post("/tests")
def create_test(data:BloodTestIn,user=Depends(roles(m.Role.DOCTOR,m.Role.CLINIC,m.Role.ADMIN)),db:Session=Depends(get_db)):
    # Authorization first: never let an unauthorized caller learn whether a
    # laboratory or accession number exists.
    assert_patient_access(db,user,data.patient_id)
    if db.scalar(select(m.BloodTest.id).where(m.BloodTest.accession_number == data.accession_number)):
        raise HTTPException(409,detail={"code":"DUPLICATE_TEST","message":"מספר הבדיקה כבר קיים","details":{}})
    if not db.get(m.Laboratory,data.laboratory_id): raise HTTPException(422,detail={"code":"UNKNOWN_LABORATORY","message":"המעבדה אינה קיימת","details":{}})
    if data.test_date>date.today(): raise HTTPException(422,detail={"code":"FUTURE_DATE","message":"תאריך הבדיקה אינו יכול להיות בעתיד","details":{}})
    test=m.BloodTest(patient_id=data.patient_id,laboratory_id=data.laboratory_id,test_date=data.test_date,accession_number=data.accession_number,source=data.source,panel=data.panel,notes=data.notes,fasting=data.fasting,reason=data.reason,created_by=user.id)
    db.add(test); db.flush()
    seen=set()
    for item in data.results:
        analyte=db.scalar(select(m.Analyte).where(m.Analyte.code==item.analyte_code.upper()))
        if not analyte: raise HTTPException(422,detail={"code":"UNKNOWN_ANALYTE","message":f"מדד לא מוכר: {item.analyte_code}","details":{}})
        if analyte.id in seen: raise HTTPException(422,detail={"code":"DUPLICATE_ANALYTE","message":"מדד מופיע פעמיים","details":{"code":analyte.code}})
        seen.add(analyte.id); db.add(m.TestResult(blood_test_id=test.id,analyte_id=analyte.id,numeric_value=item.numeric_value,text_value=item.text_value,unit=item.unit,reference_min=item.reference_min,reference_max=item.reference_max,critical_low=item.critical_low,critical_high=item.critical_high,lab_flag=item.lab_flag,data_quality_status=m.Quality.VERIFIED if item.verified else m.Quality.REQUIRES_VERIFICATION))
    AuditService.record(db,user.id,"BLOOD_TEST_CREATED","blood_test",test.id); db.commit(); db.refresh(test); run=ClinicalAnalysisEngine.analyze(db,test); AuditService.record(db,user.id,"ANALYSIS_EXECUTED","blood_test",test.id,{"algorithm_version":run.algorithm_version}); db.commit(); return {"id":test.id,"analysis_run_id":run.id}
def test_payload(t):
    return {"id":t.id,"patient_id":t.patient_id,"test_date":t.test_date,"accession_number":t.accession_number,"panel":t.panel,"source":t.source,"results":[{"id":r.id,"code":r.analyte.code,"name_he":r.analyte.display_name_he,"description_he":r.analyte.description_he,"value":r.numeric_value,"unit":r.unit,"normalized_value":r.normalized_value,"normalized_unit":r.normalized_unit,"reference_min":r.reference_min,"reference_max":r.reference_max,"reference_source":r.reference_source,"status":r.status.value,"quality":r.data_quality_status.value} for r in t.results]}
@router.get("/tests/{test_id}")
def get_test(test_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    t=authorized_test(db,user,test_id)
    return test_payload(t)
@router.get("/tests/{test_id}/analysis")
def analysis(test_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    authorized_test(db,user,test_id)
    t=db.get(m.BloodTest,test_id); run=db.scalar(select(m.AnalysisRun).where(m.AnalysisRun.blood_test_id==test_id).order_by(m.AnalysisRun.created_at.desc()))
    return {"algorithm_version":run.algorithm_version,"summary":run.summary,"quality_score":run.quality_score,"results":test_payload(t)["results"],"recommendations":RecommendationEngine.for_test(t),"disclaimer":"HEMORA היא מערכת תומכת מידע ואינה מהווה אבחנה רפואית, ייעוץ רפואי או תחליף לבדיקה ולהחלטה של איש מקצוע רפואי. טווחי ייחוס עשויים להשתנות בין מעבדות ובין מטופלים."}
@router.get("/tests/{test_id}/comparison")
def comparison(test_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    authorized_test(db,user,test_id)
    current=db.get(m.BloodTest,test_id); previous=db.scalar(select(m.BloodTest).where(m.BloodTest.patient_id==current.patient_id,m.BloodTest.test_date<current.test_date).order_by(m.BloodTest.test_date.desc()))
    if not previous: return {"previous_test":None,"message":"זוהי הבדיקה הראשונה במערכת ולכן עדיין אין בסיס להשוואה.","items":[]}
    items=TestComparisonEngine.compare(current,previous); counts={k:sum(x["trend"]==k for x in items) for k in {x["trend"] for x in items}}
    return {"previous_test":{"id":previous.id,"date":previous.test_date},"current_date":current.test_date,"items":items,"summary":counts}
@router.get("/tests/{test_id}/results/{result_id}/explanation")
def result_explanation(test_id:str,result_id:int,user=Depends(current_user),db:Session=Depends(get_db)):
    """Why HEMORA classified one result the way it did - built from stored data."""
    test=authorized_test(db,user,test_id)
    result=next((r for r in test.results if r.id==result_id),None)
    if not result: raise HTTPException(404,detail={"code":"NOT_FOUND","message":"התוצאה לא נמצאה","details":{}})
    return ExplanationService.build(db,test,result)

@router.get("/patients/{patient_id}/trends")
def trends(patient_id:str,code:str="HGB",user=Depends(current_user),db:Session=Depends(get_db)):
    assert_patient_access(db,user,patient_id)
    tests=db.scalars(select(m.BloodTest).where(m.BloodTest.patient_id==patient_id)).all(); return LongitudinalTrendService.calculate(tests,code.upper())
@router.post("/import/preview")
async def import_preview(file:UploadFile=File(...),user=Depends(roles(m.Role.CLINIC,m.Role.DOCTOR,m.Role.ADMIN))):
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024: raise HTTPException(413,"הקובץ גדול מדי")
    if (file.filename or "").lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        import csv, io
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            output = io.StringIO(); writer = csv.writer(output)
            for row in workbook.active.iter_rows(values_only=True): writer.writerow(row)
            workbook.close(); content = output.getvalue().encode("utf-8")
        except Exception: raise HTTPException(422,"קובץ Excel אינו תקין")
    elif not (file.filename or "").lower().endswith(".csv"): raise HTTPException(415,"יש לבחור CSV או XLSX")
    try: return ImportService.preview_csv(content)
    except (UnicodeError, ValueError): raise HTTPException(422,"הקובץ אינו תקין; יש לשמור CSV בקידוד UTF-8")
@router.get("/summary")
def dashboard_summary(user=Depends(current_user),db:Session=Depends(get_db)):
    """Aggregates for the dashboard, counted over the patients this role may see."""
    visible=scope_patients(select(m.Patient.id),user).subquery()
    patients=db.scalar(select(func.count()).select_from(visible)) or 0
    test_ids=select(m.BloodTest.id).where(m.BloodTest.patient_id.in_(select(visible)))
    tests=db.scalar(select(func.count()).select_from(test_ids.subquery())) or 0
    open_alerts=db.scalar(select(func.count()).select_from(
        select(m.Alert.id).where(m.Alert.patient_id.in_(select(visible)),m.Alert.state!="RESOLVED").subquery())) or 0
    abnormal=db.scalar(select(func.count()).select_from(
        select(m.TestResult.id).where(m.TestResult.blood_test_id.in_(test_ids),
            m.TestResult.status.in_([m.ResultStatus.LOW,m.ResultStatus.HIGH,m.ResultStatus.CRITICAL_LOW,m.ResultStatus.CRITICAL_HIGH])).subquery())) or 0
    quality=db.scalar(select(func.avg(m.AnalysisRun.quality_score)).where(m.AnalysisRun.blood_test_id.in_(test_ids)))
    return {"patients":patients,"tests":tests,"open_alerts":open_alerts,"abnormal_results":abnormal,
            "average_quality":round(quality,1) if quality is not None else None,
            "quality_note":"איכות הנתונים מתארת עד כמה הנתונים שהוזנו מלאים ומתאימים לניתוח ואינה ציון בריאות."}

@router.get("/clinics")
def clinics(user=Depends(current_user),db:Session=Depends(get_db)):
    rows=db.scalars(select(m.Clinic).where(m.Clinic.active.is_(True)).order_by(m.Clinic.name)).all()
    if user.role in {m.Role.DOCTOR,m.Role.CLINIC}: rows=[c for c in rows if c.id==user.clinic_id]
    return [{"id":c.id,"name":c.name,"code":c.code,"city":c.city} for c in rows]

@router.get("/patients/{patient_id}/doctors")
def patient_doctors(patient_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    assert_patient_access(db,user,patient_id)
    rows=db.scalars(select(m.DoctorPatient).where(m.DoctorPatient.patient_id==patient_id)).all()
    out=[]
    for row in rows:
        doctor=db.get(m.User,row.doctor_user_id)
        if doctor: out.append({"doctor_user_id":doctor.id,"email":doctor.email,"assigned_at":row.assigned_at})
    return out

@router.post("/patients/{patient_id}/doctors")
def assign_doctor(patient_id:str,payload:DoctorAssignment,user=Depends(roles(m.Role.ADMIN,m.Role.CLINIC)),db:Session=Depends(get_db)):
    patient=assert_patient_access(db,user,patient_id)
    doctor=db.get(m.User,payload.doctor_user_id)
    if not doctor or doctor.role!=m.Role.DOCTOR: raise HTTPException(422,detail={"code":"NOT_A_DOCTOR","message":"המשתמש שנבחר אינו רופא","details":{}})
    if user.role==m.Role.CLINIC and doctor.clinic_id!=patient.clinic_id: raise HTTPException(403,detail={"code":"FORBIDDEN","message":"ניתן לשייך רק רופא מאותה מרפאה","details":{}})
    if db.scalar(select(m.DoctorPatient.id).where(m.DoctorPatient.doctor_user_id==doctor.id,m.DoctorPatient.patient_id==patient_id)):
        raise HTTPException(409,detail={"code":"ALREADY_ASSIGNED","message":"הרופא כבר משויך למטופל זה","details":{}})
    db.add(m.DoctorPatient(doctor_user_id=doctor.id,patient_id=patient_id,assigned_by=user.id))
    AuditService.record(db,user.id,"DOCTOR_ASSIGNED","patient",patient_id,{"before":None,"after":doctor.email}); db.commit()
    return {"doctor_user_id":doctor.id,"email":doctor.email}

@router.delete("/patients/{patient_id}/doctors/{doctor_user_id}")
def unassign_doctor(patient_id:str,doctor_user_id:int,user=Depends(roles(m.Role.ADMIN,m.Role.CLINIC)),db:Session=Depends(get_db)):
    assert_patient_access(db,user,patient_id)
    row=db.scalar(select(m.DoctorPatient).where(m.DoctorPatient.doctor_user_id==doctor_user_id,m.DoctorPatient.patient_id==patient_id))
    if not row: raise HTTPException(404,detail={"code":"NOT_FOUND","message":"השיוך לא נמצא","details":{}})
    doctor=db.get(m.User,doctor_user_id); db.delete(row)
    AuditService.record(db,user.id,"DOCTOR_UNASSIGNED","patient",patient_id,{"before":doctor.email if doctor else None,"after":None}); db.commit()
    return {"removed":True}

@router.get("/analytes")
def analytes(user=Depends(current_user),db:Session=Depends(get_db)): return [{"code":a.code,"display_name_he":a.display_name_he,"category":a.category,"unit":a.typical_unit,"description_he":a.description_he} for a in db.scalars(select(m.Analyte).where(m.Analyte.active.is_(True))).all()]
@router.get("/audit")
def audit(page:int=1,user=Depends(roles(m.Role.ADMIN)),db:Session=Depends(get_db)):
    logs=db.scalars(select(m.AuditLog).order_by(m.AuditLog.timestamp.desc()).offset((page-1)*50).limit(50)).all(); return [{"id":x.id,"user_id":x.user_id,"action":x.action,"entity_type":x.entity_type,"entity_id":x.entity_id,"timestamp":x.timestamp,"metadata":x.metadata_sanitized} for x in logs]
@router.get("/reports/{test_id}.csv")
def report_csv(test_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    authorized_test(db,user,test_id)
    t=db.get(m.BloodTest,test_id); lines=["code,name_he,value,unit,status,reference_min,reference_max"]+[f'{r.analyte.code},{r.analyte.display_name_he},{r.numeric_value},{r.unit},{r.status.value},{r.reference_min or ""},{r.reference_max or ""}' for r in t.results]; AuditService.record(db,user.id,"REPORT_DOWNLOADED","blood_test",test_id,{"format":"csv"}); db.commit(); return Response("\ufeff"+"\n".join(lines),media_type="text/csv",headers={"Content-Disposition":f'attachment; filename="hemora-{test_id}.csv"'})

@router.get("/reports/{test_id}.pdf")
def report_pdf(test_id:str,user=Depends(current_user),db:Session=Depends(get_db)):
    t=authorized_test(db,user,test_id)
    content=ReportService.generate_pdf(db,t)
    AuditService.record(db,user.id,"REPORT_DOWNLOADED","blood_test",test_id,{"format":"pdf"}); db.commit(); return Response(content,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="hemora-{test_id}.pdf"'})
